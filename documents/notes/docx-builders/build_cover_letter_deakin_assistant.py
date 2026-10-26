"""Cover letter: Deakin Student Experience ASSISTANT, Burwood. HEW 3.

Different letter from build_cover_letter_deakin_sea.py. Do not send that one for this
posting; the questions are not the same.

THE POSTING ASKS THREE QUESTIONS. All answered explicitly and in order:
  1. Confirm which campus.          -> Paragraph 2. Burwood (3 vacancies).
  2. Confirm enrolment through Trimesters 1 and 2, 2027.  -> Paragraph 2.
  3. Leadership experience, and how he would use it to lead, support and motivate the
     Student Experience Ambassadors.  -> The rest, and it is most of the letter.

THE ARGUMENT, and the reason this letter works at all: Manit has no leadership TITLE
worth pitching. He was a committee member, not an exec. So the letter refuses the title
frame in its third paragraph and replaces it with the thing he can actually evidence
three times over: in every role he has held, the work ended with someone else able to do
it without him. New 7-Eleven starters on POS and food safety. New BSSG volunteers taking
over the monthly publication outright. Thirty workshop attendees debugging faults he
personally wrote and then coached them through without taking the keyboard.

That is a stronger answer than a title, and it is true. Do not "fix" it by inflating him
to club exec or by claiming he has led an ambassador team already. He has not.

The four concrete commitments near the end are the part a HEW 3 panel will remember,
because most student applicants answer this question with adjectives. Keep them specific.
They are all deducible from what he has actually done, not invented capability claims.

Honesty notes carried from README.md: committee MEMBER at DSEC; he built the workshop app
and solely designed the 15 seeded defects but did NOT deliver the session; the Open Day
ambassador role is one he has been SELECTED for, not one he has performed yet, so the
letter says "will have done the shift", future tense. Do not shift that to past tense.

House style: no em dashes, Australian spelling.

Run:  python Scripts/build_cover_letter_deakin_assistant.py
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches, Cm

NAVY = RGBColor(0x2F, 0x54, 0x96)
BLACK = RGBColor(0x00, 0x00, 0x00)
FONT = "Calibri"

OUT = (r"C:\Users\Manit Khera\Documents\Resume\Cover_Letters"
       r"\Manit_Khera_Cover_Letter_Deakin_Assistant.docx")


def bottom_rule(paragraph, color="2F5496", size=6):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def para(doc, text="", size=10.5, bold=False, color=BLACK, align=None,
         space_after=6, space_before=0, spacing=1.0, letter_spacing=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = spacing
    if text:
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        if letter_spacing is not None:
            rpr = run._element.get_or_add_rPr()
            spc = OxmlElement("w:spacing")
            spc.set(qn("w:val"), str(int(letter_spacing * 20)))
            rpr.append(spc)
    return p


def commitment(doc, lead, rest):
    """A bulleted commitment. Lead clause bold so the four scan in a quick read."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.6)
    pf.first_line_indent = Cm(-0.35)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.0
    for text, bold in (("•  ", False), (lead, True), (rest, False)):
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = Pt(10.5)
        run.font.bold = bold
        run.font.color.rgb = BLACK
    return p


OPENING = [
    "My name is Manit Khera and I am studying a Bachelor of Software Engineering "
    "(Honours) at Deakin University. I am writing to apply for the Student Experience "
    "Assistant role.",

    # Q1 and Q2, answered flatly and early.
    "To confirm your first two questions directly: I am applying for the Burwood "
    "campus. I am currently enrolled at Deakin University and I will remain enrolled, "
    "continuing my studies through Trimesters 1 and 2 of 2027. My degree runs through to "
    "2028, so I am available for the full six-month contract and any extension.",

    # Q3 begins. The title frame is refused deliberately.
    "On leadership, I will give you the honest version rather than the impressive one. I "
    "have not been a club president. What I can show you is that in every role I have "
    "held, the work ended with someone else able to do it without me.",

    "At 7-Eleven I trained new starters on the POS, food safety and shift routines while "
    "running shifts solo, opening and closing the store unsupervised. Training someone on "
    "a till during peak trade teaches you quickly that nobody learns from being talked "
    "through the whole process at once. They learn by doing one part of it while you "
    "stand close enough to catch it and far enough back that the customer is still "
    "theirs.",

    "At the Boroondara Stroke Support Group I ran the monthly activity calendar "
    "publication for 12 months, then trained new volunteers to run it themselves. The "
    "whole point was that it should stop depending on me.",

    "On the Deakin Software Engineering Club committee I built the application behind a "
    "debugging workshop and designed the 15 defects the 30 attendees had to find, then "
    "floated the room as a helper. Having written every bug myself, I knew where people "
    "would get stuck. The job was to get them unstuck without reaching over and taking "
    "the keyboard, which is harder than it sounds and is most of what supporting a team "
    "actually is. I also run the Data Science Club's public channel end to end as its "
    "LinkedIn Manager, and host and anchor our networking sessions.",

    "Here is how I would use that with the Student Experience Ambassadors.",
]

COMMITMENTS = [
    ("Brief before the shift, not during it.",
     " Most nerves on a stall are not about confidence, they are about not knowing what "
     "you are going to be asked. Ten minutes at the start on the five questions we will "
     "genuinely get that day is worth more than any amount of encouragement."),
    ("Take the first hard question, then hand the second one back.",
     " Answering everything for someone keeps them dependent on you all day. Modelling it "
     "once and then stepping back is how they finish the shift able to handle it alone."),
    ("Watch for the third hour.",
     " Orientation shifts are long and there is a predictable point where people go quiet "
     "and start letting students walk past. Rotating someone off the front for twenty "
     "minutes is not coddling them, it is how you get a good fourth hour."),
    ("Name what went well, out loud, at the end.",
     " Casual staff on a 100-hour-a-year contract rarely hear specifically what they did "
     "well. It costs nothing and it is the thing most likely to bring someone back for "
     "the next intake."),
]

CLOSING = [
    "I have also been selected as a Student Ambassador with Deakin Future Students for "
    "Open Day, which is the role this position leads. By the time this contract starts I "
    "will have worked the shift myself, and I would rather lead a team having stood where "
    "they are standing.",

    "Thank you for considering my application. I am at Burwood regularly and available "
    "across the key intake periods.",
]


def build():
    doc = Document()

    # Tighter than the Ambassador letter's geometry. This letter answers three questions
    # instead of two and carries the four commitments, so it needs the room to stay on
    # one page. Re-check the page count if you edit the body.
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)

    para(doc, "MANIT KHERA", size=20, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1, letter_spacing=2.0)
    rule_p = para(doc, "0412 761 206  |  manitkhera26@gmail.com  |  Blackburn South, "
                       "VIC 3130  |  linkedin.com/in/manit-khera-603957252",
                  size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    bottom_rule(rule_p)

    para(doc, "12 August 2026", size=10, space_before=4, space_after=7)

    for line in ["DeakinTALENT Employer Services Team", "Deakin University"]:
        para(doc, line, size=10, space_after=0, spacing=1.0)

    para(doc, "Re: Application, Student Experience Assistant, Burwood",
         size=10, bold=True, color=NAVY, space_before=9, space_after=7)

    para(doc, "Dear DeakinTALENT Employer Services Team,", size=10, space_after=6)

    for block in OPENING:
        para(doc, block, size=10, space_after=6, spacing=1.0)
    for lead, rest in COMMITMENTS:
        commitment(doc, lead, rest)
    for block in CLOSING:
        para(doc, block, size=10, space_after=6, space_before=2, spacing=1.0)

    para(doc, "Yours sincerely,", size=10, space_before=3, space_after=12)
    para(doc, "Manit Khera", size=10.5, bold=True, color=NAVY, space_after=0)

    doc.save(OUT)
    print("Written:", OUT)


if __name__ == "__main__":
    build()
