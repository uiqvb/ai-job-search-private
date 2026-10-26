"""Build the Mill Software cover letter as a .docx styled to match
1_Manit_Khera_Resume_final.pdf (Calibri, navy letterhead, rule under header).

Regenerate with:  python build_cover_letter.py
Edit the PARAGRAPHS list below to change the letter body.
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

NAVY = RGBColor(0x2F, 0x54, 0x96)
BLACK = RGBColor(0x00, 0x00, 0x00)
FONT = "Calibri"

OUT = r"C:\Users\Manit Khera\Documents\Resume\Cover_Letters\Manit_Khera_Cover_Letter_MillSoftware.docx"


def bottom_rule(paragraph, color="2F5496", size=6):
    """Draw a horizontal rule beneath a paragraph."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def para(doc, text="", size=10.5, bold=False, color=BLACK, align=None,
         space_after=6, space_before=0, spacing=1.0, caps=False, letter_spacing=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = spacing
    if text:
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.all_caps = caps
        if letter_spacing is not None:
            rpr = run._element.get_or_add_rPr()
            spc = OxmlElement("w:spacing")
            spc.set(qn("w:val"), str(int(letter_spacing * 20)))
            rpr.append(spc)
    return p


BODY = [
    "I am writing to ask to be considered for an internship or part-time software "
    "engineering position with Mill Software, having read on your careers page that you "
    "regularly take on new staff across your core service disciplines. I am a Software "
    "Engineering (Honours) student at Deakin University, based locally in Blackburn South, "
    "and I would value the opportunity to begin my career with your team.",

    "What drew me to Mill Software is the depth of the work itself. Spending time across "
    "your site, the NetTester suite stands out as something genuinely difficult and "
    "genuinely useful \u2014 Warrior Probe detecting faults and degradation across a network "
    "footprint before customers ever feel them, Defender Probe qualifying triple-play "
    "services across fibre, satellite and wireless, and InSite turning site commissioning "
    "from a paper-based process into a guided one that measurably reduces failure rates and "
    "revisits. That NetTester Web delivers carrier-grade measurement with IPv6 and NAT64 "
    "support tells me a great deal about the engineering standard behind it. Building tools "
    "that operators trust enough to run their rollouts on is a high bar, and fifteen years "
    "of doing it across telecommunications, consumer electronics and defence is not "
    "something many firms can claim.",

    "I would be applying to learn rather than to arrive finished. My interest sits exactly "
    "where your core competency does \u2014 embedded systems and Linux on ARM-based devices, "
    "and the backend services that sit above them \u2014 and the chance to develop that under "
    "engineers who build it professionally would expand my knowledge far beyond what "
    "university projects can offer.",

    "By way of background, my coursework has given me a working foundation to build on: C "
    "and C++ on microcontrollers with I\u00b2C sensor interfacing, a Raspberry Pi service "
    "deployed under systemd with MQTT telemetry, and backend development in C# and ASP.NET "
    "Core with PostgreSQL. These are student-scale projects, but they reflect where my "
    "interest genuinely lies and the standard I hold myself to \u2014 I achieved 100 in "
    "Backend Development, 100 in Data Structures and Algorithms, and 93 in Embedded "
    "Systems, and one project was selected for showcase at Deakin Open Day.",

    "In short, I offer a solid technical grounding, a strong work ethic, and real enthusiasm "
    "for the domain Mill Software works in \u2014 and I am eager to be taught. I would be "
    "grateful for the chance to contribute in whatever capacity is useful to your team. "
    "Thank you sincerely for taking the time to consider my application, and I look forward "
    "to the possibility of speaking with you at interview.",
]


def build():
    doc = Document()

    # Page setup - match a standard one-page letter
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)

    # --- Letterhead (mirrors the resume header) ---
    para(doc, "MANIT KHERA", size=24, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, letter_spacing=2.0)
    para(doc, "0412 761 206  |  manitkhera26@gmail.com", size=9.5,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    rule_p = para(doc, "Blackburn South, VIC  |  linkedin.com/in/manit-khera-603957252",
                  size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    bottom_rule(rule_p)

    # --- Date ---
    para(doc, "26 July 2026", size=10.5, space_before=6, space_after=10)

    # --- Addressee block ---
    for line in ["Hiring Team", "Mill Software", "PO Box 4545",
                 "Central Park VIC 3145"]:
        para(doc, line, size=10.5, space_after=0, spacing=1.05)

    # --- Subject line ---
    para(doc, "Re: Application \u2014 Software Engineering (Embedded / Backend)",
         size=10.5, bold=True, color=NAVY, space_before=12, space_after=10)

    # --- Salutation + body ---
    para(doc, "Dear Hiring Team,", size=10.5, space_after=8)
    for block in BODY:
        para(doc, block, size=10.5, space_after=8, spacing=1.06)

    # --- Sign-off ---
    para(doc, "Yours sincerely,", size=10.5, space_before=4, space_after=16)
    para(doc, "Manit Khera", size=11, bold=True, color=NAVY, space_after=0)

    doc.save(OUT)
    print("Written:", OUT)


if __name__ == "__main__":
    build()
