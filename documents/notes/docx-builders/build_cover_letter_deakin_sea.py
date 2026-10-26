"""Cover letter: Deakin Student Experience AMBASSADOR, Burwood. Closes 12 Aug 2026.

Styled to match the other letters in Cover_Letters/ (Calibri, navy letterhead, rule).

THE POSTING ASKS TWO QUESTIONS. Both are answered explicitly and in order:
  1. Confirm currently enrolled and remaining enrolled through Trimesters 1 and 2, 2027.
     -> Paragraph 2, stated flatly. This is a hard eligibility gate; do not bury it in
        prose or leave the reader inferring it.
  2. Why this role, and what makes him a strong and standout candidate.
     -> The rest of the letter.

THE PITCH: he has already done the recruiting half of this job for a year. Fronting
Market Day stalls for two clubs is the same work as an Orientation stall, and he has the
100+ sign-ups to show for it. The Open Day ambassador selection with Future Students is
the differentiator and it goes in early.

DELIBERATELY NOT WRITTEN:
  - No invented memory of his own Orientation. He has never described one, and a
    fabricated "a volunteer in a yellow shirt helped me find my class" is exactly the
    detail an interviewer asks him to expand on. The letter stays on what he did, not on
    a manufactured origin story. ⚠️ If Manit HAS a real memory of starting at Deakin, it
    belongs in paragraph 4 and would improve the letter. Ask him.
  - No "I have always been passionate about student experience". Same reason the Suncorp
    letter refuses the insurance version of that sentence.
  - No claim to have used Deakin support services himself. Unverified. The campus
    knowledge claim rests on two club committees and the club inbox instead, which is
    documented.

Honesty notes: Market Day 100+ sign-ups is confirmed for both DSEC and DDSC (README).
DSEC role is committee MEMBER; he did not run the workshop, he built the app behind it
and floated the room. Do not upgrade either.

House style: no em dashes, Australian spelling.

Run:  python Scripts/build_cover_letter_deakin_sea.py
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

NAVY = RGBColor(0x2F, 0x54, 0x96)
BLACK = RGBColor(0x00, 0x00, 0x00)
FONT = "Calibri"

OUT = (r"C:\Users\Manit Khera\Documents\Resume\Cover_Letters"
       r"\Manit_Khera_Cover_Letter_Deakin_SEA.docx")


def bottom_rule(paragraph, color="2F5496", size=6):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def para(doc, text="", size=10.5, bold=False, color=BLACK, align=None,
         space_after=6, space_before=0, spacing=1.0, letter_spacing=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = spacing
    if text:
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        if letter_spacing is not None:
            rpr = run._element.get_or_add_rPr()
            spc = OxmlElement("w:spacing")
            spc.set(qn("w:val"), str(int(letter_spacing * 20)))
            rpr.append(spc)
    return p


BODY = [
    "My name is Manit Khera and I am studying a Bachelor of Software Engineering "
    "(Honours) at Deakin University, Burwood. I am writing to apply for the Student "
    "Experience Ambassador role at the Burwood campus.",

    # Q1, answered flatly and early.
    "To confirm your first question directly: I am currently enrolled at Deakin "
    "University and I will remain enrolled, continuing my studies through Trimesters 1 "
    "and 2 of 2027. My degree runs through to 2028, so I am available for the full "
    "six-month contract and for the extension if it is offered.",

    "I am applying because I have spent the past year doing an unofficial version of "
    "this job and I would like to do the real one. I have been selected as a Student "
    "Ambassador with Deakin Future Students for Open Day. Before that, I fronted the "
    "Market Day stall for two Deakin clubs, the Software Engineering Club and the Data "
    "Science Club, which meant standing at a table for hours having the same first "
    "conversation over and over with students who had no idea who we were, and signing "
    "up over 100 new members between them.",

    "That is the part of this role I actually want. Orientation is a room full of people "
    "who do not know anyone yet, and it matters enormously who the first person to speak "
    "to them is. I have watched students walk past our stall three times before working "
    "up to stopping, and I have learned that the useful move is almost never a better "
    "pitch. It is asking them what they are studying and letting them talk.",

    "On what I would bring beyond enthusiasm, three things. I have hosted and anchored "
    "networking sessions for the Data Science Club, connecting current students with "
    "alumni and industry guests, so I am comfortable holding a microphone in front of a "
    "room and keeping a conversation moving between people who have just met. I floated "
    "a 30-attendee club workshop as a helper, which is mostly the skill of scanning a "
    "room for the person who has stopped typing and is too embarrassed to put their hand "
    "up. And I know the Burwood campus and Deakin's services from two years here, two "
    "club committees, and managing a club inbox where students asked us questions we "
    "often had to redirect to the right part of the University.",

    "Casual work at peak intake also needs people who reliably turn up and hold the "
    "unglamorous end of the job. For the past 12 months I have run the monthly activity "
    "calendar publication for the Boroondara Stroke Support Group, a community "
    "not-for-profit, publishing to a deadline every single month and then training new "
    "volunteers to do it themselves. At 7-Eleven I worked the front line of a "
    "high-volume store, opened and closed unsupervised, and took escalated complaints "
    "through to resolution without a manager. Neither is glamorous, and both taught me "
    "that composure under pressure is mostly preparation plus not taking it personally.",

    "I am at Burwood regularly, I am available across the key intake periods, and I "
    "would like the chance to do this across a year rather than a single Open Day.",

    "Thank you for considering my application.",
]


def build():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)

    para(doc, "MANIT KHERA", size=24, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, letter_spacing=2.0)
    para(doc, "0412 761 206  |  manitkhera26@gmail.com", size=9.5,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    rule_p = para(doc, "Blackburn South, VIC 3130  |  "
                       "linkedin.com/in/manit-khera-603957252",
                  size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    bottom_rule(rule_p)

    para(doc, "12 August 2026", size=10.5, space_before=6, space_after=10)

    for line in ["DeakinTALENT Employer Services Team", "Deakin University"]:
        para(doc, line, size=10.5, space_after=0, spacing=1.05)

    para(doc, "Re: Application, Student Experience Ambassador, Burwood",
         size=10.5, bold=True, color=NAVY, space_before=12, space_after=10)

    para(doc, "Dear DeakinTALENT Employer Services Team,", size=10.5, space_after=8)
    for block in BODY:
        para(doc, block, size=10.5, space_after=8, spacing=1.06)

    para(doc, "Yours sincerely,", size=10.5, space_before=4, space_after=16)
    para(doc, "Manit Khera", size=11, bold=True, color=NAVY, space_after=0)

    doc.save(OUT)
    print("Written:", OUT)


if __name__ == "__main__":
    build()
