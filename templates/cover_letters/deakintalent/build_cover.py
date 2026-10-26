"""Build a cover letter from the DeakinTALENT Word template.

Reads a plain-text draft (<file>.txt), fills template.docx beside this script,
saves <file>.docx next to the draft, then exports <file>.pdf through Microsoft Word
and prints the page count. Word does the layout, so the PDF matches what the
employer sees if they open the .docx.

Draft format: sections headed by "## <Name>", in this order. Blank lines separate
paragraphs inside "Body"; every other section is one item per line.

    ## Sender
    Your Name
    1 Example Street
    Suburb, VIC 3000
    04XX XXX XXX
    you@example.com

    ## Date
    Tuesday 20 October 2026

    ## Recipient
    Hiring Manager
    Organisation Name
    Street Address
    Suburb, State Postcode

    ## Greeting
    Dear Hiring Manager,

    ## RE
    RE: Software Engineering Intern position, REF: 12345   (leave the section empty to drop it)

    ## Body
    Introduction ...

    Body paragraph one ...

    ## Sign-off
    Kind regards,

    ## Name
    Your Name

**text** inside a paragraph becomes bold. The last two Sender lines (phone, email)
form the second right-aligned block, as in the template.

Run from the draft's folder:
    python ../templates/cover_letters/deakintalent/build_cover.py cover_<company>_<role>.txt
"""

import copy
import re
import sys
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "template.docx"
SECTIONS = ["Sender", "Date", "Recipient", "Greeting", "RE", "Body", "Sign-off", "Name"]
BOLD = re.compile(r"(\*\*[^*]+\*\*)")


def parse(path):
    text = path.read_text(encoding="utf-8")
    parts = re.split(r"^##\s*(.+?)\s*$", text, flags=re.M)
    found = {name.strip(): body.strip() for name, body in zip(parts[1::2], parts[2::2])}
    missing = [s for s in SECTIONS if s not in found]
    if missing:
        sys.exit(f"Draft is missing section(s): {', '.join(missing)}")
    lines = lambda s: [l.strip() for l in found[s].splitlines() if l.strip()]
    body = [" ".join(l.strip() for l in b.splitlines()) for b in re.split(r"\n\s*\n", found["Body"]) if b.strip()]
    sender = lines("Sender")
    if len(sender) < 3:
        sys.exit("Sender needs at least name, address and the phone/email lines")
    return dict(sender=sender[:-2], contact=sender[-2:], date=" ".join(lines("Date")),
                recipient=lines("Recipient"), greeting=" ".join(lines("Greeting")),
                re_line=" ".join(lines("RE")), body=body,
                signoff=" ".join(lines("Sign-off")), name=" ".join(lines("Name")))


def fill(p, text):
    """Put text into a one-run placeholder paragraph, honouring **bold** spans."""
    proto = p.runs[0]
    base_bold = proto.bold
    segs = [s for s in BOLD.split(text) if s]
    anchor = proto._element
    for _ in segs[1:]:
        new = copy.deepcopy(proto._element)
        anchor.addnext(new)
        anchor = new
    for run, seg in zip(p.runs, segs):
        is_bold = seg.startswith("**") and seg.endswith("**")
        run.text = seg[2:-2] if is_bold else seg
        run.bold = True if is_bold else base_bold


def expand(doc, token, items, separator=False):
    """Replace the placeholder paragraph with one cloned paragraph per item."""
    ph = next(p for p in doc.paragraphs if p.text == token)
    blank = next(p for p in doc.paragraphs if not p.text)
    anchor = ph._element
    for i, item in enumerate(items):
        if separator and i:
            b = copy.deepcopy(blank._element)
            anchor.addnext(b)
            anchor = b
        el = copy.deepcopy(ph._element)
        anchor.addnext(el)
        anchor = el
        fill(next(p for p in doc.paragraphs if p._element is el), item)
    ph._element.getparent().remove(ph._element)


def drop_with_following_blank(doc, token):
    ph = next(p for p in doc.paragraphs if p.text == token)
    nxt = ph._element.getnext()
    ph._element.getparent().remove(ph._element)
    nxt.getparent().remove(nxt)


def export_pdf(docx_path):
    import win32com.client as win32
    word = win32.gencache.EnsureDispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(str(docx_path), ReadOnly=True)
        try:
            pages = doc.ComputeStatistics(2)
            words = doc.ComputeStatistics(0)
            doc.SaveAs(str(docx_path.with_suffix(".pdf")), FileFormat=17)
        finally:
            doc.Close(SaveChanges=False)
    finally:
        word.Quit()
    return pages, words


def main(draft):
    draft = Path(draft).resolve()
    L = parse(draft)
    doc = Document(str(TEMPLATE))
    # The sender block is two right-aligned paragraphs joined by line breaks, as in the original.
    for token, items in (("[SENDER_LINE]", L["sender"]), ("[SENDER_CONTACT_LINE]", L["contact"])):
        p = next(p for p in doc.paragraphs if p.text == token)
        p.runs[0].text = items[0]
        for item in items[1:]:
            p.runs[-1].add_break()
            r = copy.deepcopy(p.runs[0]._element)
            p._p.append(r)
            p.runs[-1].text = item
    expand(doc, "[DATE]", [L["date"]])
    expand(doc, "[RECIPIENT_LINE]", L["recipient"])
    expand(doc, "[GREETING]", [L["greeting"]])
    if L["re_line"]:
        expand(doc, "[RE_LINE]", [L["re_line"]])
    else:
        drop_with_following_blank(doc, "[RE_LINE]")
    expand(doc, "[BODY_PARAGRAPH]", L["body"], separator=True)
    expand(doc, "[SIGN_OFF]", [L["signoff"]])
    expand(doc, "[YOUR_NAME]", [L["name"]])
    doc.core_properties.author = L["name"]
    doc.core_properties.title = "Cover Letter"

    out = draft.with_suffix(".docx")
    try:
        doc.save(out)
    except PermissionError:
        sys.exit(f"Cannot write {out.name}: close it in Word and re-run")
    pdf = out.with_suffix(".pdf")
    pdf.unlink(missing_ok=True)  # never leave a stale PDF behind a failed export
    pages, words = export_pdf(out)
    print(f"Written: {out.name}, {pdf.name} ({pages} page(s), {words} words)")
    if pages != 1:
        sys.exit(f"FAIL: cover letter is {pages} pages, must be exactly 1")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("usage: build_cover.py <draft>.txt")
    main(sys.argv[1])
